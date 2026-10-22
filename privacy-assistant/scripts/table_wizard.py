"""Local-only wizard worker. Its metadata/preview files never go to the Agent."""
import csv,hashlib,hmac,io,json,math,os,sys,unicodedata
from pathlib import Path
from zipfile import ZipFile
ROOT=Path(__file__).resolve().parents[1]
CORE=Path(r'C:\PrivacySkill') if sys.platform=='win32' and Path(r'C:\PrivacySkill').exists() else ROOT/'skill'
sys.path.insert(0,str(CORE/'scripts'));sys.path.insert(0,str(ROOT/'runtime/Lib/site-packages'))
import deidentify as d


def raw_rows(path):
    data=d.read_bytes(path,d.MAX_FILE)
    if path.suffix.lower()=='.csv':
        rows=list(csv.reader(io.StringIO(data.decode('utf-8-sig')),strict=True))
    elif path.suffix.lower()=='.xlsx':
        with ZipFile(io.BytesIO(data)) as z:
            d.require(sum(x.file_size for x in z.infolist())<=d.MAX_EXPANDED,'resource_limit')
            d.require(not any(any(k in x.filename.lower() for k in ('vbaproject','externallinks','embeddings')) for x in z.infolist()),'active_content')
        import openpyxl
        w=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=False)
        try:
            d.require(w.worksheets[0].max_row<=d.MAX_ROWS+1 and w.worksheets[0].max_column<=d.MAX_COLS,'resource_limit')
            rows=[]
            for line in w.worksheets[0].iter_rows():
                d.require(not any(c.data_type=='f' for c in line),'active_content');rows.append([c.value for c in line])
        finally:w.close()
    else:raise d.Block('input_rejected')
    d.require(rows and len(rows)<=d.MAX_ROWS+1 and 1<=len(rows[0])<=d.MAX_COLS,'schema_mismatch')
    headers=rows[0]
    d.require(all(isinstance(h,str) and h.strip() for h in headers) and len(set(headers))==len(headers),'schema_mismatch')
    d.require(all(len(r)==len(headers) for r in rows),'schema_mismatch')
    return rows


def inspect(task):
    source=source_path(task);rows=raw_rows(source)
    fields=[]
    for i,h in enumerate(rows[0]):fields.append({'index':i,'header':h,'kind':'id' if h.strip() in ('学号','工号','编号') else 'drop'})
    # Ambiguous IDs must be chosen locally, never silently combined.
    if sum(f['kind']=='id' for f in fields)!=1:
        for f in fields:f['kind']='drop'
    write_private(task,'schema.json',{'fields':fields,'first_sheet_only':True})


def source_path(task):
    candidates=[p for p in (task/'private').glob('original.*') if p.suffix.lower() in ('.csv','.xlsx')]
    d.require(len(candidates)==1,'input_rejected');return candidates[0]


def write_private(task,name,data):
    private=task/'private';d.safe_path(private);private.mkdir(exist_ok=True,mode=0o700)
    dest=private/name;d.safe_path(dest)
    dest.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8');os.chmod(dest,0o600)


def configure(task):
    source=source_path(task);rows=raw_rows(source)
    choices=json.loads(d.read_bytes(task/'private/choices.json',128*1024))
    d.require(isinstance(choices,list) and len(choices)==len(rows[0]),'policy_invalid')
    columns=[];quasi=[]
    for i,(header,kind) in enumerate(zip(rows[0],choices)):
        d.require(kind in ('id','drop','number','category'),'policy_invalid');c={'header':header,'kind':kind}
        # A novice must not accidentally select a phone/card number as a measure.
        normalized=unicodedata.normalize('NFKC',header).lower().replace(' ','')
        direct=('姓名','学号','工号','手机号','手机号码','电话','身份证','证件号','银行卡','邮箱','email','phone','身份证明','家庭住址','微信号','qq号')
        if any(s in normalized for s in direct):d.require(kind in ('id','drop'),'policy_invalid')
        if kind=='id':c['output']='person'
        if kind=='number':c.update(output='measure_'+str(i+1),min=-1e12,max=1e12,missing=['','NA'])
        if kind=='category':
            vals=sorted({r[i] for r in rows[1:] if isinstance(r[i],str)})
            d.require(1<=len(vals)<=50 and all(isinstance(r[i],str) for r in rows[1:]),'category_unknown')
            c.update(output='category_'+str(i+1),mapping={v:'value_'+str(n+1) for n,v in enumerate(vals)});quasi.append(c['output'])
        columns.append(c)
    d.require(sum(c['kind']=='id' for c in columns)==1,'policy_invalid')
    write_private(task,'policy.json',{'version':1,'tables':[{'sheet_index':0,'output':'students','columns':columns,'quasi_identifiers':quasi,'min_group':5}]})
    d.load_policy(task/'private/policy.json')

def column_guide(task):
    # Only a fixed vocabulary crosses the boundary, never arbitrary raw headers.
    allowed={'分数','得分','总分','年龄','年级','性别','学院','班级','人数','次数','时长','完成情况'}
    table=d.load_policy(task/'private/policy.json')['tables'][0]
    return [{'output':c['output'],'kind':c['kind'],'meaning':
             ('身份代号' if c['kind']=='id' else c['header'].strip() if c['header'].strip() in allowed else '数值' if c['kind']=='number' else '分类')}
            for c in table['columns'] if c['kind']!='drop']


def preview(task):
    p=d.load_policy(task/'private/policy.json');t=p['tables'][0];cols=[c for c in t['columns'] if c['kind']!='drop']
    source=source_path(task);job=task/'job';binding=json.loads(d.read_bytes(job/'private/job.json',1024*1024))
    d.require(hashlib.sha256(d.read_bytes(source,d.MAX_FILE)).hexdigest()==binding['input_sha256'],'job_mismatch')
    d.require(hashlib.sha256(json.dumps(p,sort_keys=True).encode()).hexdigest()==binding['policy_sha256'],'job_mismatch')
    data=d.read_bytes(task/'private/returned.csv',d.MAX_FILE)
    header=next(csv.reader(io.StringIO(data.decode('utf-8-sig'))))
    names=[c['output'] for c in cols];d.require(header[:len(names)]==names,'result_mismatch')
    extra=header[len(names):];d.require(extra in ([],['follow_up'],['priority'],['follow_up','priority']),'result_mismatch')
    derived=[]
    if 'follow_up' in extra:derived.append({'output':'follow_up','kind':'category','mapping':{'需跟进':'yes','无需跟进':'no'}})
    if 'priority' in extra:derived.append({'output':'priority','kind':'number','min':0,'max':3,'nullable':False})
    result=d.result_rows(task/'private/returned.csv',cols+derived);d.require(bool(result),'result_mismatch')
    baseline=job/'private/students.csv'
    d.require(hashlib.sha256(d.read_bytes(baseline,d.MAX_FILE)).hexdigest()==binding['tables']['students'],'job_mismatch')
    base={r['person']:r for r in d.result_rows(baseline,cols)}
    for r in result:
        d.require(r['person'] in base,'result_mismatch')
        d.require(all(r[c['output']]==base[r['person']][c['output']] for c in cols),'unauthorized_change')
    original,_,_=d.input_tables(source,p);raw=original[0][1:]
    key=d.read_bytes(job/'private/job-key.bin',64)
    index=next(i for i,c in enumerate(t['columns']) if c['kind']=='id')
    def token(row):return 'p_'+hmac.new(key,unicodedata.normalize('NFKC',row[index]).strip().encode(),hashlib.sha256).hexdigest()[:32]
    by_token={token(r):r for r in raw};d.require(set(base)==set(by_token),'result_mismatch')
    fields=[{'source':'original','column':index,'header':t['columns'][index]['header']}]
    name_index=next((i for i,c in enumerate(t['columns']) if c['header'].strip()=='姓名'),None)
    if name_index is not None:fields.append({'source':'original','column':name_index,'header':'姓名'})
    # Include the retained numeric/category columns and any bounded new results.
    for i,c in enumerate(t['columns']):
        if c['kind'] in ('number','category'):fields.append({'source':'original','column':i,'header':c['header']})
    for name in extra:fields.append({'source':'result','column':name,'header':'是否跟进' if name=='follow_up' else '优先级'})
    q={'table':'students','editable':[],'derived':derived,'selection':{'mode':'subset','tokens':[r['person'] for r in result]},'fields':fields}
    # Snapshot bytes before review, so the Agent cannot change the reviewed CSV.
    (task/'private/reviewed.csv').write_bytes(data)
    write_private(task,'import-policy.json',q)
    display=[]
    for r in result:
        original_row=by_token[r['person']];x={}
        for f in fields:
            v=original_row[f['column']] if f['source']=='original' else r[f['column']]
            if f['column']=='follow_up':v={'yes':'需跟进','no':'无需跟进'}[v]
            x[f['header']]=None if v is None else str(v)
        display.append(x)
    write_private(task,'preview.json',display)


def main():
    try:
        os.umask(0o077);mode=sys.argv[1];task=Path(sys.argv[2]);d.isolation_check();d.safe_path(task)
        if mode=='inspect':inspect(task)
        elif mode=='configure':configure(task)
        elif mode=='preview':preview(task)
        else:raise d.Block('invalid_mode')
        print(json.dumps({'status':'local_ready','cloud_release':False}));return 0
    except d.Block as e:print(json.dumps({'status':'blocked','code':str(e),'cloud_release':False}));return 2
    except Exception:print(json.dumps({'status':'blocked','code':'processing_failed','cloud_release':False}));return 2

if __name__=='__main__':sys.exit(main())
