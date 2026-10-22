"""Synthetic regression coverage for the local launcher and table wizard."""
import csv,hashlib,io,json,os,shutil,subprocess,sys,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
ROOT=Path(__file__).resolve().parents[1]/'privacy-assistant'
sys.path.insert(0,str(ROOT/'scripts'))
import assistant as a
import table_wizard as w
import openpyxl

class Wizard(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(prefix='privacy-wizard-fake-');self.task=Path(self.temp.name);(self.task/'private').mkdir()
        self.raw=self.task/'private/original.csv'
        self.rows=[['学号','姓名','分组','分数','备注']]+[[f'{i:06d}',f'虚构学生{i:03d}','虚构甲组' if i%2==0 else '虚构乙组',20+i%17,'PRIVATE_TEST_SENTINEL'] for i in range(60)]
        with self.raw.open('w',newline='') as f:csv.writer(f).writerows(self.rows)
    def tearDown(self):self.temp.cleanup()
    def worker(self,mode,isolated=True):
        cmd=[sys.executable,'-B',str(ROOT/'scripts/table_wizard.py'),mode,str(self.task)]
        if isolated:cmd=['/usr/bin/sandbox-exec','-p','(version 1)(allow default)(deny network*)']+cmd
        r=subprocess.run(cmd,capture_output=True,text=True,timeout=30)
        self.assertNotIn('PRIVATE_TEST_SENTINEL',r.stdout+r.stderr);self.assertNotIn('虚构学生',r.stdout+r.stderr);self.assertNotIn(str(self.task),r.stdout+r.stderr)
        return r.returncode,json.loads(r.stdout)
    def configure(self,choices=None):
        (self.task/'private/choices.json').write_text(json.dumps(choices or ['id','drop','category','number','drop']))
        return self.worker('configure')
    def core(self,extra):
        return subprocess.run([sys.executable,'-B',str(ROOT/'skill/scripts/run_local.py'),'--input',str(self.raw),'--policy',str(self.task/'private/policy.json')]+extra,capture_output=True,text=True,timeout=40)
    def prepare(self):
        self.assertEqual(self.configure()[0],0)
        r=self.core(['--out',str(self.task/'job')]);self.assertEqual(r.returncode,0,r.stdout)
        with (self.task/'job/private/students.csv').open() as f:reader=csv.DictReader(f);self.headers=reader.fieldnames;self.result=list(reader)
    def returned(self,rows=None,extra=None):
        with (self.task/'private/returned.csv').open('w',newline='') as f:
            writer=csv.DictWriter(f,fieldnames=self.headers+(extra or []));writer.writeheader();writer.writerows(self.result if rows is None else rows)
    def test_default_all_non_id_dropped(self):
        self.assertEqual(self.worker('inspect')[0],0)
        fields=json.loads((self.task/'private/schema.json').read_text())['fields']
        self.assertEqual([x['kind'] for x in fields],['id','drop','drop','drop','drop'])
    def test_worker_rejects_unisolated_before_output(self):
        code,r=self.worker('inspect',False);self.assertNotEqual(code,0);self.assertEqual(r['code'],'isolation_failed');self.assertFalse((self.task/'private/schema.json').exists())
    def test_ambiguous_ids_require_local_choice(self):
        self.rows[0][1]='工号'
        with self.raw.open('w',newline='') as f:csv.writer(f).writerows(self.rows)
        self.assertEqual(self.worker('inspect')[0],0)
        self.assertTrue(all(x['kind']=='drop' for x in json.loads((self.task/'private/schema.json').read_text())['fields']))
    def test_two_id_columns_rejected(self):self.assertNotEqual(self.configure(['id','id','drop','number','drop'])[0],0)
    def test_unknown_choice_rejected(self):self.assertNotEqual(self.configure(['id','text','drop','number','drop'])[0],0)
    def test_phone_misclassified_as_measure_rejected(self):
        self.rows[0][3]='手机号'
        with self.raw.open('w',newline='') as f:csv.writer(f).writerows(self.rows)
        self.assertNotEqual(self.configure()[0],0)
    def test_id_only_supported(self):self.assertEqual(self.configure(['id','drop','drop','drop','drop'])[0],0)
    def test_full_roundtrip(self):
        self.prepare();self.result.reverse();self.returned();self.assertEqual(self.worker('preview')[0],0)
        r=self.core(['--mode','prepare-import','--job',str(self.task/'job'),'--results',str(self.task/'private/reviewed.csv'),'--import-policy',str(self.task/'private/import-policy.json'),'--out',str(self.task/'export')]);self.assertEqual(r.returncode,0,r.stdout)
        wb=openpyxl.load_workbook(self.task/'export/private/ready-to-import.xlsx');actual=list(wb.active.values);wb.close()
        expected=[tuple(str(v) for v in x[:4]) for x in self.rows];self.assertEqual(actual,expected)
    def test_subset_new_columns_exact_restoration(self):
        self.prepare();selected=[dict(r,follow_up='yes',priority='2') for r in self.result if int(r['measure_4'])>=30];selected.reverse()
        self.returned(selected,['follow_up','priority']);self.assertEqual(self.worker('preview')[0],0)
        reviewed=(self.task/'private/reviewed.csv').read_bytes();(self.task/'private/returned.csv').write_text('changed after review')
        self.assertEqual(reviewed,(self.task/'private/reviewed.csv').read_bytes())
        r=self.core(['--mode','prepare-import','--job',str(self.task/'job'),'--results',str(self.task/'private/reviewed.csv'),'--import-policy',str(self.task/'private/import-policy.json'),'--out',str(self.task/'export')]);self.assertEqual(r.returncode,0,r.stdout)
        wb=openpyxl.load_workbook(self.task/'export/private/ready-to-import.xlsx');actual=list(wb.active.values)[1:];wb.close()
        expected=[tuple([str(v) for v in x[:4]]+['需跟进',2]) for x in self.rows[1:] if x[3]>=30]
        self.assertEqual(len(expected),21);self.assertEqual(actual,expected)
    def test_preview_rejects_original_column_change(self):
        self.prepare();self.result[0]['measure_4']='99';self.returned();code,r=self.worker('preview');self.assertNotEqual(code,0);self.assertEqual(r['code'],'unauthorized_change')
    def test_preview_rejects_duplicates(self):
        self.prepare();self.returned([self.result[0],self.result[0]]);self.assertEqual(self.worker('preview')[1]['code'],'duplicate_identity')
    def test_preview_rejects_unknown_identity(self):
        self.prepare();self.result[0]['person']='p_'+'0'*32;self.returned();self.assertEqual(self.worker('preview')[1]['code'],'result_mismatch')
    def test_preview_rejects_empty_subset(self):
        self.prepare();self.returned([]);self.assertEqual(self.worker('preview')[1]['code'],'result_mismatch')
    def test_new_column_out_of_range(self):
        self.prepare();self.returned([dict(r,priority='100') for r in self.result],['priority']);self.assertEqual(self.worker('preview')[1]['code'],'numeric_invalid')
    def test_unknown_new_column(self):
        self.prepare();self.returned([dict(r,notes='hello') for r in self.result],['notes']);self.assertEqual(self.worker('preview')[1]['code'],'result_mismatch')
    def test_baseline_tampering_rejected(self):
        self.prepare();self.returned();(self.task/'job/private/students.csv').write_text('tampered');self.assertEqual(self.worker('preview')[1]['code'],'job_mismatch')
    def test_source_tampering_rejected(self):
        self.prepare();self.returned();self.raw.write_text('tampered');self.assertEqual(self.worker('preview')[1]['code'],'job_mismatch')
    def test_formula_in_dropped_field_rejected(self):
        self.rows[1][-1]='=HYPERLINK("https://invalid")'
        with self.raw.open('w',newline='') as f:csv.writer(f).writerows(self.rows)
        self.assertEqual(self.configure()[0],0);r=self.core(['--out',str(self.task/'job')]);self.assertEqual(json.loads(r.stdout)['code'],'active_content')
    def test_xlsx_numeric_type_preserved(self):
        self.raw.unlink();self.raw=self.raw.with_suffix('.xlsx')
        wb=openpyxl.Workbook()
        for row in self.rows:wb.active.append(row)
        wb.save(self.raw);wb.close()
        self.prepare();self.returned();self.assertEqual(self.worker('preview')[0],0)
        r=self.core(['--mode','prepare-import','--job',str(self.task/'job'),'--results',str(self.task/'private/reviewed.csv'),'--import-policy',str(self.task/'private/import-policy.json'),'--out',str(self.task/'export')]);self.assertEqual(r.returncode,0,r.stdout)
        wb=openpyxl.load_workbook(self.task/'export/private/ready-to-import.xlsx');actual=list(wb.active.values);wb.close()
        self.assertEqual(actual,[tuple(x[:4]) for x in self.rows])

class Package(unittest.TestCase):
    def test_two_measure_correlation_on_mac_python(self):
        from decimal import Decimal
        rows=[{'x':Decimal(x),'y':Decimal(y)} for x,y in [(1,1),(2,3),(3,2)]]
        result=w.d.stats(rows,[{'kind':'number','output':'x'},{'kind':'number','output':'y'}])
        self.assertAlmostEqual(result['correlations']['x__y']['r'],0.5,places=14)
    def test_release_integrity(self):self.assertEqual(a.check_package()['version'],'0.6.1')
    def test_no_personal_paths_in_authored_files(self):
        for folder in ('scripts','skills'):
            for p in (ROOT/folder).rglob('*'):
                if p.is_file() and p.suffix in ('.py','.ps1','.md'):
                    self.assertNotRegex(p.read_text(),r'/Users/[A-Za-z][^/\s]*/')
    def test_installed_does_not_claim_isolation(self):
        r=a.doctor();self.assertEqual(r['status'],'setup_ready');self.assertFalse(r['current_agent_isolated']);self.assertFalse(r['windows_device_validated'])
    def test_cloud_platform_blocked(self):
        with patch.object(a.sys,'platform','linux'):self.assertEqual(a.doctor()['code'],'local_desktop_required')
    def test_mac_workbuddy_not_silently_launched(self):
        with tempfile.TemporaryDirectory() as t:self.assertEqual(a.start('workbuddy',False,Path(t))['code'],'mac_workbuddy_gui_not_isolated')
    def test_install_idempotent_keeps_unrelated_skill(self):
        with tempfile.TemporaryDirectory() as t:
            base=Path(t)/'state';home=Path(t)/'skills';home.mkdir();(home/'other.md').write_text('preserve')
            for _ in range(2):self.assertEqual(a.install('codex',base,home)['status'],'installed')
            self.assertEqual((home/'other.md').read_text(),'preserve');self.assertTrue((home/'privacy-assistant/SKILL.md').exists())
    def test_unmanaged_skill_not_overwritten(self):
        with tempfile.TemporaryDirectory() as t:
            base=Path(t)/'state';home=Path(t)/'skills';p=home/'privacy-assistant';p.mkdir(parents=True);(p/'SKILL.md').write_text('mine')
            with self.assertRaisesRegex(ValueError,'existing_skill_preserved'):a.install('codex',base,home)
            self.assertEqual((p/'SKILL.md').read_text(),'mine')
    def test_windows_encoded_commands_fit_limit(self):
        import base64
        for name in ('launcher.ps1','wizard_ui.ps1'):self.assertLess(len(base64.b64encode((ROOT/'scripts'/name).read_text().encode('utf-16le')))+500,32767)
    def test_windows_online_mapping_never_contains_vault(self):
        from windows_agent import configuration
        import xml.etree.ElementTree as ET
        root=ET.fromstring(configuration(Path('C:/Exchange/task'),Path('C:/bundle')))
        self.assertEqual([x.text for x in root.findall('MappedFolders/MappedFolder/SandboxFolder')],['C:\\Exchange','C:\\AgentKit']);self.assertEqual(root.findtext('ClipboardRedirection'),'Disable')
    def test_windows_offline_configuration_has_no_network(self):
        from windows_host import configuration
        import xml.etree.ElementTree as ET
        root=ET.fromstring(configuration([(Path('C:/task'),'C:\\PrivacyTask',False)]));self.assertEqual(root.findtext('Networking'),'Disable');self.assertEqual(root.findtext('ClipboardRedirection'),'Disable')

if __name__=='__main__':unittest.main(verbosity=2)
