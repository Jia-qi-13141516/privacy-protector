"""Background-only regressions. Real model acceptance is recorded separately."""
import io,json,os,subprocess,sys,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
ROOT=Path(__file__).resolve().parents[1]/'privacy-assistant'
sys.path.insert(0,str(ROOT/'scripts'))
import table_wizard
import background_engine as e
import background as bg
import document_engine as doc
import task_channel as ch
import openpyxl

class Background(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.base=Path(self.temp.name)
        self.source=self.base/'fictional.xlsx';book=openpyxl.Workbook();sheet=book.active
        self.rows=[['编号','姓名','待补材料数','已处理份数'],['0001','虚构甲',2,9],['0002','虚构乙',3,8],['0003','虚构丙',5,7],['0004','虚构丁',None,6]]
        for row in self.rows:sheet.append(row)
        book.save(self.source);book.close()
        self.task=self.base/'private';self.task.mkdir()
    def tearDown(self):self.temp.cleanup()
    def run_table(self,prompt):return e.table_job(self.source,self.task,prompt,None)
    def output(self):
        book=openpyxl.load_workbook(self.task/'result.xlsx');values=list(book['处理结果'].values);book.close();return values
    def test_business_meaning_and_full_values_preserved(self):
        before=self.source.read_bytes();r=self.run_table('筛选待补材料数不低于3，新增是否跟进为需跟进、优先级统一为1，不要修改原值。')
        self.assertEqual(r['rows'],2)
        self.assertEqual(self.output(),[tuple(self.rows[0]+['是否跟进','优先级']),tuple(self.rows[2]+['需跟进',1]),tuple(self.rows[3]+['需跟进',1])])
        self.assertEqual(before,self.source.read_bytes());self.assertEqual(r['local_model_calls'],0)
    def test_zero_match_is_valid_empty_workbook(self):
        r=self.run_table('筛选待补材料数大于99，并统计记录数');self.assertTrue(r['zero_match']);self.assertEqual(self.output(),[tuple(self.rows[0])])
    def test_sort_keeps_identity_bound_to_its_values(self):
        self.run_table('筛选待补材料数>=3，并按待补材料数从大到小排列')
        self.assertEqual(self.output()[1:], [tuple(self.rows[3]),tuple(self.rows[2])])
    def test_completed_and_role_only_are_not_todos(self):
        source=self.base/'f.txt';source.write_text('顾衡是宋岚的项目联络人，宋岚负责物料。\n顾衡尚未交材料，需要在下周一交给宋岚。\n陶芷已经完成核对，该事项已结束，无需再次跟进。\n韩舟负责海报，已发布归档，不再安排相关工作。')
        class Model:
            calls=0
            def ask(self,*args):
                self.calls+=1
                if self.calls==1:return {'terms':['顾衡','宋岚','陶芷']}
                offered=json.loads(args[1])['段落'];self.offered=[p['id'] for p in offered]
                return {'selected_ids':self.offered}
        model=Model();e.document_job(source,self.task,'提取未完成待办',model)
        self.assertEqual(model.offered,[2])
        self.assertIn('韩舟',json.loads((self.task/'identity-map.json').read_text()).values())
    def test_non_numeric_sort_is_not_silently_unicode_collation(self):
        with self.assertRaisesRegex(ValueError,'task_not_supported'):self.run_table('按姓名降序排序')
    def test_missing_value_not_converted_to_zero(self):
        self.run_table('筛选待补材料数小于3，并统计记录数');self.assertEqual(self.output()[1:], [tuple(self.rows[1])])
    def test_two_numeric_conditions_use_actual_columns(self):
        self.run_table('筛选待补材料数>=3且已处理份数>=8，统计记录数');self.assertEqual(self.output()[1:],[tuple(self.rows[2])])
    def test_unknown_column_rejected_not_silently_dropped(self):
        with self.assertRaisesRegex(ValueError,'task_not_supported'):self.run_table('筛选不存在的业务数>=3，统计记录数')
        self.assertFalse((self.task/'result.xlsx').exists())
    def test_unexpressed_range_rejected(self):
        with self.assertRaisesRegex(ValueError,'task_not_supported'):self.run_table('筛选待补材料数>=3且<5，统计记录数')
    def test_or_not_silently_changed_to_and(self):
        with self.assertRaisesRegex(ValueError,'task_not_supported'):self.run_table('筛选待补材料数>=3或已处理份数>=8，统计记录数')
    def test_unsupported_sort_not_reported_completed(self):
        with self.assertRaisesRegex(ValueError,'task_not_supported'):self.run_table('筛选待补材料数>=3，按姓名和编号降序排序')
    def test_unsupported_grouping_not_reported_completed(self):
        with self.assertRaisesRegex(ValueError,'task_not_supported'):self.run_table('按每个学院统计人数')
    def test_arbitrary_new_column_not_ignored(self):
        with self.assertRaisesRegex(ValueError,'task_not_supported'):self.run_table('新增评价备注列，统计记录数')
    def test_formula_in_csv_header_rejected(self):
        source=self.base/'fictional.csv';source.write_text('=HYPERLINK("https://example.invalid"),数量\n假数据,3\n')
        with self.assertRaises(ValueError):e.table_job(source,self.task,'统计记录数',None)
    def test_result_column_collision_stops(self):
        book=openpyxl.load_workbook(self.source);book.active.cell(1,5,'是否跟进');book.save(self.source);book.close()
        with self.assertRaisesRegex(ValueError,'result_column_exists'):self.run_table('统计记录数，新增是否跟进为需跟进')
    def test_unicode_numeric_precision_preserved(self):
        self.rows[1][2]=3.125;book=openpyxl.Workbook()
        for r in self.rows:book.active.append(r)
        book.save(self.source);book.close();self.run_table('筛选待补材料数>3.12，统计记录数')
        self.assertEqual(self.output()[1][2],3.125)
    def test_no_source_read_without_os_isolation(self):
        if sys.platform!='darwin':self.skipTest('Mac actual OS probe')
        result=subprocess.run([sys.executable,'-B',str(ROOT/'scripts/background_worker.py'),str(self.base/'does-not-exist')],capture_output=True,text=True)
        self.assertEqual(result.returncode,2);self.assertEqual(json.loads(result.stdout)['code'],'isolation_failed')
        self.assertNotIn(str(self.source),result.stdout+result.stderr)
    def test_kernel_denies_internet_and_loopback_in_child(self):
        if sys.platform!='darwin':self.skipTest('Mac actual OS boundary')
        code='import sys;sys.path.insert(0,sys.argv[1]);from background_worker import require_offline;require_offline();print("PASS")'
        r=subprocess.run(['/usr/bin/sandbox-exec','-p','(version 1)(allow default)(deny network*)(deny mach-lookup)(deny appleevent-send)',sys.executable,'-B','-c',code,str(ROOT/'scripts')],capture_output=True,text=True)
        self.assertEqual(r.returncode,0);self.assertEqual(r.stdout.strip(),'PASS')
    def test_public_error_never_contains_raw_exception(self):
        self.assertEqual(bg.safe_error(ValueError('姓名：PRIVATE_CANARY; /private/path')),'local_task_failed')
    def test_task_id_cannot_escape_private_queue(self):
        with self.assertRaises(ValueError):bg.state(self.base,'../../private')
    def test_document_unknown_segment_cannot_be_restored(self):
        source=self.base/'f.txt';source.write_text('林安须在周五前把材料交给周宁。')
        class Model:
            calls=0
            def ask(self,*args):
                self.calls+=1;return {'terms':['林安','周宁']} if self.calls==1 else {'selected_ids':[999]}
        with self.assertRaisesRegex(ValueError,'local_inference_invalid'):e.document_job(source,self.task,'整理待办事项',Model())
        self.assertFalse((self.task/'result.docx').exists())
    def test_document_output_keeps_exact_relationships(self):
        source=self.base/'f.txt';source.write_text('林安须把材料交给周宁，不是周宁交给林安。\n陈禾已全部完成，没有后续任务。')
        class Model:
            calls=0
            def ask(self,*args):
                self.calls+=1;return {'terms':['林安','周宁','陈禾']} if self.calls==1 else {'selected_ids':[1]}
        result=e.document_job(source,self.task,'整理待办事项',Model())
        self.assertEqual(result['paragraphs'],1)
        output=doc.docx_text((self.task/'result.docx').read_bytes())
        self.assertIn('林安须把材料交给周宁，不是周宁交给林安。',output);self.assertNotIn('陈禾',output)
    def test_documents_never_claim_arbitrary_rewrite(self):
        with self.assertRaisesRegex(ValueError,'task_not_supported'):e.document_job(self.base/'unused.txt',self.task,'改写这份文档',None)

if __name__=='__main__':unittest.main()
