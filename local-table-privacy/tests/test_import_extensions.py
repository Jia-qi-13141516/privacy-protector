"""All fixtures are synthetic; test actual process isolation and output cells."""
import csv,json,copy,sys,tempfile,unittest,subprocess
from pathlib import Path
import openpyxl
from test_privacy import fixture,RUN,SENTINEL

class ImportExtensions(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.d=Path(self.tmp.name)
        self.input,self.policy,_=fixture(self.d);self.job=self.d/'job';self.out=self.d/'import'
        self.assertEqual(self.run_cmd(['--out',str(self.job)])['status'],'local_ready')
        with (self.job/'private/wave_0.csv').open() as f:self.base=list(csv.DictReader(f))
        self.rows=copy.deepcopy(self.base)
        self.q={'table':'wave_0','editable':[], 'derived':[
          {'output':'follow_up','kind':'category','mapping':{'需跟进':'yes','无需跟进':'no'}},
          {'output':'priority','kind':'number','min':0,'max':3,'nullable':False}],
          'selection':{'mode':'all'},'fields':[
          {'source':'original','column':0,'header':'学号'}, {'source':'original','column':1,'header':'姓名'},
          {'source':'result','column':'follow_up','header':'是否跟进'}, {'source':'result','column':'priority','header':'优先级'}]}
        self.headers=list(self.base[0])+['follow_up','priority']
        for r in self.rows:r.update(follow_up='yes' if int(r['score'])>=30 else 'no',priority='2')
        self.results=self.d/'results.csv';self.qp=self.d/'import.json'
    def tearDown(self):self.tmp.cleanup()
    def run_cmd(self,args):
        r=subprocess.run([sys.executable,'-B',str(RUN),'--input',str(self.input),'--policy',str(self.policy)]+args,capture_output=True,text=True)
        for s in (SENTINEL,'虚构学生',str(self.d)):self.assertNotIn(s,r.stdout+r.stderr)
        return json.loads(r.stdout)
    def call(self):
        self.qp.write_text(json.dumps(self.q))
        with self.results.open('w',newline='') as f:
            w=csv.DictWriter(f,fieldnames=self.headers);w.writeheader();w.writerows(self.rows)
        return self.run_cmd(['--out',str(self.out),'--mode','prepare-import','--job',str(self.job),'--results',str(self.results),'--import-policy',str(self.qp)])
    def blocked(self,code):self.assertEqual(self.call()['code'],code);self.assertFalse(self.out.exists())
    def select(self):
        self.rows=[r for r in self.rows if r['follow_up']=='yes']
        self.q['selection']={'mode':'subset','tokens':[r['person'] for r in self.rows]}
    def test_new_columns_exact_values_and_source_order(self):
        self.rows.reverse();self.assertEqual(self.call()['status'],'local_ready')
        w=openpyxl.load_workbook(self.out/'private/ready-to-import.xlsx');rows=list(w.active.values);w.close()
        self.assertEqual(rows[0],('学号','姓名','是否跟进','优先级'))
        for i,r in enumerate(rows[1:]):self.assertEqual(r,(f'{i:06d}',f'虚构学生{i:03d}','需跟进' if 20+i%17>=30 else '无需跟进',2))
    def test_subset_exact_independent_ids(self):
        self.select();self.rows.reverse();self.assertEqual(self.call()['status'],'local_ready')
        w=openpyxl.load_workbook(self.out/'private/ready-to-import.xlsx');rows=list(w.active.values)[1:];w.close()
        expected=[(f'{i:06d}',f'虚构学生{i:03d}','需跟进',2) for i in range(60) if 20+i%17>=30]
        self.assertEqual(rows,expected)
        receipt=json.loads((self.out/'receipt.json').read_text());self.assertEqual(receipt['exported_rows'],len(expected));self.assertEqual(receipt['omitted_rows'],60-len(expected))
    def test_subset_without_authorization(self):self.rows=self.rows[:2];self.blocked('result_mismatch')
    def test_extra_row_outside_selection(self):self.select();self.rows.append(self.base[0]|{'follow_up':'yes','priority':'1'});self.blocked('result_mismatch')
    def test_missing_selected_row(self):self.select();self.rows.pop();self.blocked('result_mismatch')
    def test_unknown_selected_token(self):self.select();self.q['selection']['tokens'][0]='p_'+'f'*32;self.blocked('result_mismatch')
    def test_duplicate_selection(self):self.select();self.q['selection']['tokens']*=2;self.blocked('policy_invalid')
    def test_empty_selection(self):self.q['selection']={'mode':'subset','tokens':[]};self.blocked('policy_invalid')
    def test_changed_protected_column(self):self.rows[0]['score']='99';self.blocked('unauthorized_change')
    def test_new_column_out_of_range(self):self.rows[0]['priority']='4';self.blocked('numeric_invalid')
    def test_new_column_formula(self):self.rows[0]['priority']='=1+1';self.blocked('numeric_invalid')
    def test_new_category_unknown(self):self.rows[0]['follow_up']='其他';self.blocked('category_unknown')
    def test_free_text_column_not_allowed(self):self.q['derived'][0]={'kind':'text','output':'follow_up'};self.blocked('policy_invalid')
    def test_missing_required_number(self):self.rows[0]['priority']='';self.blocked('numeric_invalid')
    def test_nullable_number(self):
        self.q['derived'][1]['nullable']=True;self.rows[0]['priority']='';self.assertEqual(self.call()['status'],'local_ready')
        w=openpyxl.load_workbook(self.out/'private/ready-to-import.xlsx');self.assertIsNone(w.active['D2'].value);w.close()
    def test_existing_column_collision(self):self.q['derived'][1]['output']='score';self.blocked('policy_invalid')
    def test_new_column_not_declared(self):self.q['derived'].pop();self.blocked('policy_invalid')
    def test_symlink_job(self):
        alias=self.d/'alias';alias.symlink_to(self.job,target_is_directory=True);self.job=alias;self.blocked('input_rejected')
    def test_symlink_job_parent(self):
        alias=self.d/'alias';alias.symlink_to(self.d,target_is_directory=True);self.job=alias/'job';self.blocked('input_rejected')
    def test_hardlink_results_rejected(self):
        # Write then link; call rewrites same inode and must still reject.
        self.results.touch();(self.d/'copy.csv').hardlink_to(self.results);self.blocked('input_rejected')
    def test_no_silent_merge_duplicate_result(self):self.rows.append(self.rows[0].copy());self.blocked('duplicate_identity')
    def test_header_injection_cannot_make_formula(self):
        self.q['fields'][2]['header']='=1+1';self.assertEqual(self.call()['status'],'local_ready')
        w=openpyxl.load_workbook(self.out/'private/ready-to-import.xlsx');self.assertEqual(w.active['C1'].data_type,'s');w.close()
        self.assertFalse((self.out/'private/ready-to-paste.tsv').exists())

    def test_output_ancestor_link_rejected(self):
        real=self.d/'actual';real.mkdir();alias=self.d/'alias';alias.symlink_to(real,target_is_directory=True)
        for mode in ('deidentify','prepare-summary'):
            out=alias/mode
            args=['--mode',mode,'--out',str(out)]
            if mode=='prepare-summary':
                sp=self.d/'summary.json';sp.write_text(json.dumps({'table':'wave_0','numeric':['score'],'category':None}))
                args+=['--job',str(self.job),'--summary-policy',str(sp)]
            self.assertEqual(self.run_cmd(args)['code'],'input_rejected');self.assertFalse(out.exists())

if __name__=='__main__':unittest.main()
