"""Synthetic-only continuity tests. Uses the actual process network sandbox."""
import copy
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from test_privacy import fixture, ROOT, RUN, SENTINEL

class SummaryTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.dir=Path(self.tmp.name)
        self.input,self.policy,self.spec=fixture(self.dir)
        self.job=self.dir/'job';self.out=self.dir/'summary'
        self.q={'table':'wave_0','numeric':['score'],'category':'group'}
        self.qp=self.dir/'summary-policy.json'
        r=self.run_command(['--out',str(self.job)])
        self.assertEqual(r['status'],'local_ready')
    def tearDown(self):self.tmp.cleanup()
    def run_command(self,extra):
        r=subprocess.run([sys.executable,'-B',str(RUN),'--input',str(self.input),'--policy',str(self.policy)]+extra,capture_output=True,text=True,timeout=130)
        self.assertNotIn(SENTINEL,r.stdout+r.stderr);self.assertNotIn('虚构学生',r.stdout+r.stderr);self.assertNotIn(str(self.dir),r.stdout+r.stderr)
        x=json.loads(r.stdout);self.assertFalse(x['cloud_release']);return x
    def call(self):
        self.qp.write_text(json.dumps(self.q),encoding='utf-8')
        return self.run_command(['--out',str(self.out),'--mode','prepare-summary','--job',str(self.job),'--summary-policy',str(self.qp)])
    def blocked(self,code):
        self.assertEqual(self.call()['code'],code);self.assertFalse(self.out.exists());self.assertFalse(list(self.dir.glob('.privacy-summary-*')))
    def test_summary_matches_independent_values(self):
        self.assertEqual(self.call()['code'],'summary_review_required')
        x=json.loads((self.out/'private/summary-review.json').read_text())
        self.assertEqual(x['row_count'],60);self.assertEqual(x['category_counts'],{'category_1':30,'category_2':30})
        self.assertEqual(x['numeric'],{'measure_1':{'n':60,'missing':0,'mean_rounded':'27.40'}})
        text=(self.out/'private/summary-review.json').read_text()
        for v in ['学号','姓名','手机号','虚构甲组','group_a','score','p_',SENTINEL]:self.assertNotIn(v,text)
        self.assertEqual((self.out/'private/summary-review.json').stat().st_mode&0o777,0o600)
    def test_host_loopback_stays_available(self):
        import socket
        self.call()
        with socket.socket() as server:
            server.bind(('127.0.0.1',0));server.listen(1)
            with socket.create_connection(server.getsockname(),timeout=1) as client:
                conn,_=server.accept();conn.close()
    def test_all_missing_is_null_not_zero(self):
        import openpyxl
        w=openpyxl.load_workbook(self.input)
        for ws in w.worksheets[:2]:
            for i in range(60):ws.cell(i+2,5).value=None
        w.save(self.input);w.close();self.job=self.dir/'job-empty'
        self.assertEqual(self.run_command(['--out',str(self.job)])['status'],'local_ready')
        self.call();x=json.loads((self.out/'private/summary-review.json').read_text())
        self.assertEqual(x['numeric']['measure_1'],{'n':0,'missing':60,'mean_rounded':None})
    def test_no_selected_measures_gives_only_total(self):
        self.q.update(numeric=[],category=None);self.call()
        x=json.loads((self.out/'private/summary-review.json').read_text());self.assertEqual(x['numeric'],{});self.assertEqual(x['category_counts'],{})
    def test_small_missing_count_blocks_entire_candidate(self):
        self.q['numeric']=['outcome'];self.blocked('summary_small_cell')
    def test_direct_identifier_cannot_be_measure(self):
        self.q['numeric']=['person'];self.blocked('policy_invalid')
    def test_arbitrary_text_and_filter_not_accepted(self):
        self.q['note']=SENTINEL;self.q['filter']='person';self.blocked('policy_invalid')
    def test_duplicate_measures(self):
        self.q['numeric']=['score','score'];self.blocked('policy_invalid')
    def test_bad_category(self):
        self.q['category']='person';self.blocked('policy_invalid')
    def test_changed_input(self):
        self.input.write_bytes(self.input.read_bytes()+b'changed');self.blocked('job_mismatch')
    def test_changed_baseline(self):
        p=self.job/'private/wave_0.csv';p.write_text(p.read_text()+'\n');self.blocked('job_mismatch')
    def test_existing_output_preserved(self):
        self.call();before=(self.out/'private/summary-review.json').read_bytes();self.assertEqual(self.call()['code'],'output_exists');self.assertEqual((self.out/'private/summary-review.json').read_bytes(),before)
    def test_unisolated_worker_stops_before_reading(self):
        r=subprocess.run([sys.executable,str(ROOT/'scripts/deidentify.py'),'--mode','prepare-summary','--input',str(self.dir/'missing'),'--policy',str(self.dir/'missing'),'--out',str(self.out),'--job',str(self.job),'--summary-policy',str(self.qp)],capture_output=True,text=True)
        self.assertEqual(json.loads(r.stdout)['code'],'isolation_failed');self.assertFalse(self.out.exists())
    def test_rare_category_blocks_even_when_old_minimum_passed(self):
        import openpyxl
        w=openpyxl.load_workbook(self.input)
        for ws in w.worksheets[:2]:
            for i in range(60):ws.cell(i+2,4,'虚构甲组' if i<5 else '虚构乙组')
        w.save(self.input);w.close();self.job=self.dir/'job-rare'
        self.assertEqual(self.run_command(['--out',str(self.job)])['status'],'local_ready')
        self.blocked('summary_small_cell')

if __name__=='__main__':unittest.main()
