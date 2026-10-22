#!/usr/bin/env python3
"""Strict allowlist processor. Raw input and private artifacts never go to stdout."""
import argparse
import csv
from datetime import datetime
from collections import Counter
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import hashlib
import hmac
import io
import itertools
import json
import math
import os
from pathlib import Path
import re
import secrets
import shutil
import socket
import statistics
import stat
import sys
import tempfile
import unicodedata
from zipfile import ZipFile, BadZipFile, ZipInfo, ZIP_DEFLATED
import xml.etree.ElementTree as ET

MAX_FILE = 20 * 1024 * 1024
MAX_EXPANDED = 80 * 1024 * 1024
MAX_ROWS = 50000
MAX_COLS = 100
TOKEN = re.compile(r'^[a-z][a-z0-9_]{0,31}$')


class Block(Exception):
    pass


def require(condition, code):
    if not condition:
        raise Block(code)


def isolation_check():
    # Run BEFORE reading any input. ECONNREFUSED is not isolation evidence.
    if sys.platform == 'win32':
        import subprocess
        # The OS boundary is the generated .wsb with Networking=Disable.
        # Fail closed on an ordinary host or an unexpectedly connected guest.
        require(os.environ.get('USERNAME','').lower() == 'wdagutilityaccount', 'isolation_failed')
        require(Path(sys.executable).parent == Path(r'C:\PrivacyRuntime'), 'isolation_failed')
        require(Path(__file__).resolve().parent == Path(r'C:\PrivacySkill\scripts'), 'isolation_failed')
        check = "$ErrorActionPreference='Stop'; if (@(Get-NetAdapter -IncludeHidden | Where-Object { $_.Status -eq 'Up' -and $_.InterfaceType -ne 24 }).Count -ne 0) { exit 2 }; if ((Get-CimInstance Win32_ComputerSystem).Model -ne 'Virtual Machine') { exit 3 }; [Console]::Write('NO_NETWORK_ADAPTER')"
        try:
            probe=subprocess.run([str(Path(os.environ['SystemRoot'])/'System32/WindowsPowerShell/v1.0/powershell.exe'),'-NoProfile','-NonInteractive','-Command',check],capture_output=True,text=True,timeout=20)
            require(probe.returncode == 0 and probe.stdout == 'NO_NETWORK_ADAPTER', 'isolation_failed')
        except Exception:
            raise Block('isolation_failed')
    elif sys.platform == 'darwin':
        s = socket.socket()
        try:
            s.connect(('127.0.0.1', 9))
        except PermissionError:
            pass
        except OSError:
            raise Block('isolation_failed')
        else:
            raise Block('isolation_failed')
        finally:
            s.close()
    elif sys.platform == 'linux':
        # Dedicated offline VM (no NIC), or a network namespace with only lo.
        # The marker alone is insufficient: actual kernel interfaces are checked.
        marker=Path('/etc/privacy-assistant-offline')
        require(marker.is_file() and not marker.is_symlink(), 'isolation_failed')
        info=marker.stat()
        require(info.st_uid==0 and info.st_mode & 0o022 == 0, 'isolation_failed')
        require(marker.read_text().strip()=='privacy-assistant-offline-v1', 'isolation_failed')
        require(set(p.name for p in Path('/sys/class/net').iterdir())=={'lo'}, 'isolation_failed')
        require(len(Path('/proc/net/route').read_text().splitlines())<=1, 'isolation_failed')
    else:
        raise Block('isolation_failed')
    def audit(event, args):
        if event.startswith(('socket.', 'subprocess.')) or event in ('os.system', 'os.exec', 'os.posix_spawn'):
            raise Block('isolation_failed')
    sys.addaudithook(audit)


def safe_path(path):
    # macOS system aliases are expected; user-created links/junctions are not.
    aliases = {'/tmp', '/var', '/etc'} if sys.platform == 'darwin' else set()
    for part in [path.absolute(), *path.absolute().parents]:
        if str(part) in aliases:
            continue
        require(not part.is_symlink() and not getattr(part, 'is_junction', lambda: False)(), 'input_rejected')


def read_bytes(path, limit):
    safe_path(path)
    require(path.is_file() and path.stat().st_nlink == 1, 'input_rejected')
    fd=os.open(path,os.O_RDONLY|getattr(os,'O_NOFOLLOW',0)|getattr(os,'O_BINARY',0))
    with os.fdopen(fd,'rb') as f:
        info=os.fstat(f.fileno())
        require(stat.S_ISREG(info.st_mode) and info.st_nlink==1,'input_rejected')
        data = f.read(limit + 1)
    require(len(data) <= limit, 'resource_limit')
    return data


def load_policy(path):
    try:
        p = json.loads(read_bytes(path, 1024 * 1024))
        require(isinstance(p, dict) and set(p) == {'version', 'tables'}, 'policy_invalid')
        require(p['version'] == 1 and isinstance(p['tables'], list) and 1 <= len(p['tables']) <= 2, 'policy_invalid')
        table_codes = []
        for t in p['tables']:
            require(set(t) == {'sheet_index', 'output', 'columns', 'quasi_identifiers', 'min_group'}, 'policy_invalid')
            require(type(t['sheet_index']) is int and t['sheet_index'] >= 0, 'policy_invalid')
            require(isinstance(t['output'], str) and TOKEN.fullmatch(t['output']), 'policy_invalid')
            table_codes.append(t['output'])
            require(isinstance(t['columns'], list) and 1 <= len(t['columns']) <= MAX_COLS, 'policy_invalid')
            names = []
            headers = []
            ids = 0
            for c in t['columns']:
                require(isinstance(c, dict) and {'header', 'kind'} <= set(c), 'policy_invalid')
                require(isinstance(c['header'], str) and c['header'].strip(), 'policy_invalid')
                headers.append(c['header'])
                kind = c['kind']
                allowed = {'drop': {'header','kind'}, 'id': {'header','kind','output'},
                           'number': {'header','kind','output','min','max','missing'},
                           'category': {'header','kind','output','mapping'}}
                require(kind in allowed and set(c) == allowed[kind], 'policy_invalid')
                if kind == 'drop':
                    continue
                require(isinstance(c['output'], str) and TOKEN.fullmatch(c['output']), 'policy_invalid')
                names.append(c['output'])
                if kind == 'id':
                    ids += 1
                if kind == 'number':
                    require(type(c['min']) in (int,float) and type(c['max']) in (int,float)
                            and math.isfinite(c['min']) and math.isfinite(c['max'])
                            and c['min'] <= c['max'], 'policy_invalid')
                    require(isinstance(c['missing'], list) and all(isinstance(x,str) for x in c['missing']), 'policy_invalid')
                    # Numeric values cannot be declared missing silently.
                    for x in c['missing']:
                        try:
                            Decimal(x)
                        except InvalidOperation:
                            continue
                        raise Block('policy_invalid')
                if kind == 'category':
                    m = c['mapping']
                    require(isinstance(m,dict) and 1 <= len(m) <= 50, 'policy_invalid')
                    require(all(isinstance(k,str) and isinstance(v,str) and TOKEN.fullmatch(v) for k,v in m.items()), 'policy_invalid')
                    # Do not silently collapse analytically distinct categories.
                    require(len(set(m.values())) == len(m), 'policy_invalid')
            require(ids == 1 and len(set(names)) == len(names) and len(set(headers)) == len(headers), 'policy_invalid')
            require(type(t['min_group']) is int and t['min_group'] >= 5, 'policy_invalid')
            require(isinstance(t['quasi_identifiers'], list) and len(set(t['quasi_identifiers'])) == len(t['quasi_identifiers']), 'policy_invalid')
            category_names = [c['output'] for c in t['columns'] if c['kind']=='category']
            require(set(t['quasi_identifiers']) == set(category_names), 'policy_invalid')
        require(len(set(table_codes)) == len(table_codes), 'policy_invalid')
        require(len({t['sheet_index'] for t in p['tables']}) == len(p['tables']), 'policy_invalid')
        return p
    except Block:
        raise
    except Exception:
        raise Block('policy_invalid')


def input_tables(path, policy):
    data = read_bytes(path, MAX_FILE)
    ext = path.suffix.lower()
    if ext == '.csv':
        require(len(policy['tables']) == 1 and policy['tables'][0]['sheet_index']==0, 'policy_invalid')
        try:
            content = data.decode('utf-8-sig')
            require('\x00' not in content, 'input_rejected')
            rows = list(csv.reader(io.StringIO(content), strict=True))
        except (UnicodeError, csv.Error):
            raise Block('input_rejected')
        require(len(rows) <= MAX_ROWS + 1 and all(len(r) <= MAX_COLS for r in rows), 'resource_limit')
        # CSV has no executable formulas to evaluate; reject formula-looking cells
        # even in dropped columns, allowing strictly numeric negative values later.
        for row in rows:
            for value in row:
                require(not value.lstrip().startswith(('=', '+', '@')), 'active_content')
                if value.lstrip().startswith('-'):
                    try:
                        require(Decimal(value).is_finite(), 'active_content')
                    except InvalidOperation:
                        raise Block('active_content')
        return [rows], {'hidden_sheets_removed': 0, 'comments_removed': 0, 'extra_sheets_removed': 0}, hashlib.sha256(data).hexdigest()
    require(ext == '.xlsx', 'input_rejected')
    try:
        with ZipFile(io.BytesIO(data)) as z:
            entries = z.infolist()
            require(len(entries) <= 2000 and sum(i.file_size for i in entries) <= MAX_EXPANDED, 'resource_limit')
            names = [i.filename for i in entries]
            require(len(names)==len(set(names)), 'input_rejected')
            for i in entries:
                require(not i.flag_bits & 1 and '..' not in Path(i.filename).parts and not i.filename.startswith('/'), 'input_rejected')
                require(i.file_size <= 1024*1024 or i.file_size/max(1,i.compress_size) <= 200, 'resource_limit')
                low=i.filename.lower()
                require(not any(x in low for x in ['vbaproject','externallinks/','embeddings/','connections.xml','querytables/','customxml/']), 'active_content')
                if low.endswith(('.xml','.rels')):
                    blob=z.read(i.filename)
                    require(b'\x00' not in blob, 'active_content')
                    require(b'<!DOCTYPE' not in blob.upper() and b'<!ENTITY' not in blob.upper(), 'active_content')
                    # External hyperlinks/linked images are also rejected, not fetched.
                    require(not re.search(rb'TargetMode\s*=\s*[\"\']External[\"\']',blob,re.I), 'active_content')
                    tree=ET.fromstring(blob)
                    require(not any(e.get('TargetMode','').lower()=='external' for e in tree.iter()), 'active_content')
    except BadZipFile:
        raise Block('input_rejected')
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
    comments=0
    selected={t['sheet_index'] for t in policy['tables']}
    require(all(i < len(wb.worksheets) for i in selected), 'schema_mismatch')
    require(not list(wb.defined_names), 'active_content')
    for idx,ws in enumerate(wb.worksheets):
        require(ws.max_row <= MAX_ROWS+1 and ws.max_column<=MAX_COLS, 'resource_limit')
        require(not ws.merged_cells.ranges, 'schema_mismatch')
        if idx in selected:
            require(ws.sheet_state=='visible', 'schema_mismatch')
            require(not any(x.hidden for x in ws.row_dimensions.values()) and not any(x.hidden for x in ws.column_dimensions.values()), 'schema_mismatch')
        for row in ws:
            for c in row:
                require(c.data_type not in ('f','e'), 'active_content')
                require(c.hyperlink is None, 'active_content')
                if c.comment:
                    comments+=1
    tables = [[list(r) for r in wb.worksheets[t['sheet_index']].iter_rows(values_only=True)] for t in policy['tables']]
    removed={'hidden_sheets_removed':sum(w.sheet_state!='visible' for w in wb.worksheets),
             'comments_removed':comments, 'extra_sheets_removed':len(wb.worksheets)-len(selected)}
    wb.close()
    return tables,removed,hashlib.sha256(data).hexdigest()


def number(value, c):
    if value is None or isinstance(value,str) and value in c.get('missing',[]):
        return None
    require(not isinstance(value,bool), 'numeric_invalid')
    require(isinstance(value,(str,int,float,Decimal)), 'numeric_invalid')
    if isinstance(value,str):
        # Do not silently normalize decimal separators, date values, formulas, IDs.
        require(bool(re.fullmatch(r'-?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?',value)), 'numeric_invalid')
    try:
        d=Decimal(str(value))
        require(d.is_finite() and Decimal(str(c['min']))<=d<=Decimal(str(c['max'])), 'numeric_invalid')
        require(len(d.as_tuple().digits)<=30, 'numeric_invalid')
        return d
    except InvalidOperation:
        raise Block('numeric_invalid')


def stats(rows, cols):
    result={'rows':len(rows),'numeric':{},'categories':{},'correlations':{}}
    numeric=[c for c in cols if c['kind']=='number']
    for c in numeric:
        name=c['output'];v=[r[name] for r in rows if r[name] is not None]
        total=sum(v,Decimal(0));n=len(v)
        result['numeric'][name]={'n':n,'missing':len(rows)-n,'sum':str(total),
            'mean':str(total/n) if n else None,
            'variance':str(statistics.variance(v)) if n>1 else None,
            'min':str(min(v)) if n else None,'max':str(max(v)) if n else None}
    for c in cols:
        if c['kind']=='category':result['categories'][c['output']]=dict(Counter(r[c['output']] for r in rows))
    for a,b in itertools.combinations(numeric,2):
        vals=[(float(r[a['output']]),float(r[b['output']])) for r in rows if r[a['output']] is not None and r[b['output']] is not None]
        corr=None
        if len(vals)>1:
            x,y=zip(*vals)
            if len(set(x))>1 and len(set(y))>1:corr=statistics.correlation(x,y)
        result['correlations'][a['output']+'__'+b['output']]={'n':len(vals),'r':corr}
    return result


def paired(tables, specs):
    if len(tables)!=2:return None
    keys=[next(c['output'] for c in t['columns'] if c['kind']=='id') for t in specs]
    a={r[keys[0]]:r for r in tables[0]};b={r[keys[1]]:r for r in tables[1]};shared=set(a)&set(b)
    common=set(c['output'] for c in specs[0]['columns'] if c['kind']=='number') & set(c['output'] for c in specs[1]['columns'] if c['kind']=='number')
    diffs={}
    for name in sorted(common):
        vals=[b[k][name]-a[k][name] for k in sorted(shared) if a[k][name] is not None and b[k][name] is not None]
        diffs[name]={'n':len(vals),'sum':str(sum(vals,Decimal(0))),
                     'mean':str(sum(vals,Decimal(0))/len(vals)) if vals else None}
    return {'matched':len(shared),'left_only':len(set(a)-set(b)),'right_only':len(set(b)-set(a)),'differences':diffs}


def process(input_path, policy_path, output):
    isolation_check()
    require(not output.exists() and not output.is_symlink(), 'output_exists')
    safe_path(output)
    p=load_policy(policy_path)
    raw,removed,input_digest=input_tables(input_path,p)
    key=secrets.token_bytes(32);identities={};token_reverse={};original=[];clean=[]
    for rows,t in zip(raw,p['tables']):
        require(rows and rows[0]==[c['header'] for c in t['columns']], 'schema_mismatch')
        require(len(rows)>1 and len(rows)-1>=t['min_group'], 'small_group')
        seen=set();original_rows=[];clean_rows=[]
        for row in rows[1:]:
            require(len(row)==len(t['columns']), 'schema_mismatch')
            old={};new={}
            for v,c in zip(row,t['columns']):
                kind=c['kind']
                if kind=='drop':continue
                name=c['output']
                if kind=='id':
                    require(isinstance(v,str) and bool(v.strip()), 'missing_identity')
                    norm=unicodedata.normalize('NFKC',v).strip()
                    require(1<=len(norm)<=128 and not any(unicodedata.category(x).startswith('C') for x in norm), 'ambiguous_identity')
                    require(norm not in identities or identities[norm]==v, 'ambiguous_identity')
                    identities[norm]=v
                    require(norm not in seen, 'duplicate_identity');seen.add(norm)
                    token='p_'+hmac.new(key,norm.encode(),hashlib.sha256).hexdigest()[:32]
                    require(token not in token_reverse or token_reverse[token]==norm, 'ambiguous_identity')
                    token_reverse[token]=norm
                    old[name]=norm;new[name]=token
                elif kind=='number':old[name]=new[name]=number(v,c)
                elif kind=='category':
                    require(isinstance(v,str) and v in c['mapping'], 'category_unknown')
                    old[name]=new[name]=c['mapping'][v]
            original_rows.append(old);clean_rows.append(new)
        qis=t['quasi_identifiers']
        if qis:
            counts=Counter(tuple(r[n] for n in qis) for r in clean_rows)
            require(min(counts.values())>=t['min_group'], 'small_group')
        original.append(original_rows);clean.append(clean_rows)
    # Stage privately; publish only after readback and all integrity checks pass.
    require(output.parent.is_dir() and not output.parent.is_symlink(), 'input_rejected')
    stage=Path(tempfile.mkdtemp(prefix='.privacy-',dir=output.parent));os.chmod(stage,0o700)
    try:
        private=stage/'private';private.mkdir(mode=0o700)
        reread=[]
        for records,t in zip(clean,p['tables']):
            cols=[c for c in t['columns'] if c['kind']!='drop'];names=[c['output'] for c in cols]
            dest=private/(t['output']+'.csv')
            with dest.open('w',encoding='utf-8',newline='') as f:
                writer=csv.DictWriter(f,fieldnames=names);writer.writeheader();writer.writerows(records)
            os.chmod(dest,0o600)
            with dest.open(encoding='utf-8',newline='') as f:
                read=[]
                for row in csv.DictReader(f):
                    read.append({c['output']:(None if row[c['output']]=='' else Decimal(row[c['output']])) if c['kind']=='number' else row[c['output']] for c in cols})
            require(read==records,'integrity_failed');reread.append(read)
        before=[stats(r,t['columns']) for r,t in zip(original,p['tables'])]
        after=[stats(r,t['columns']) for r,t in zip(reread,p['tables'])]
        pair_before=paired(original,p['tables']);pair_after=paired(reread,p['tables'])
        require(before==after and pair_before==pair_after,'integrity_failed')
        (private/'job-key.bin').write_bytes(key);os.chmod(private/'job-key.bin',0o600)
        require(hashlib.sha256(read_bytes(input_path,MAX_FILE)).hexdigest()==input_digest,'integrity_failed')
        require(load_policy(policy_path)==p,'integrity_failed')
        binding={'input_sha256':input_digest,
                 'policy_sha256':hashlib.sha256(json.dumps(p,sort_keys=True).encode()).hexdigest(),
                 'tables':{t['output']:hashlib.sha256((private/(t['output']+'.csv')).read_bytes()).hexdigest() for t in p['tables']}}
        (private/'job.json').write_text(json.dumps(binding),encoding='utf-8');os.chmod(private/'job.json',0o600)
        # No original -> pseudonym mapping is written. The key still remains sensitive.
        metrics={'before':before,'after':after,'paired_before':pair_before,'paired_after':pair_after}
        (private/'statistics.json').write_text(json.dumps(metrics,ensure_ascii=False,indent=2),encoding='utf-8')
        os.chmod(private/'statistics.json',0o600)
        receipt={'status':'local_ready','code':'local_ready','cloud_release':False,
                 'network_denied':True,'integrity_passed':True,
                 'checks':['cell_values','row_order','missingness','sum_mean_variance','pairwise_correlation','category_counts','cross_table_matching'],
                 'anonymity_certified':False,'metadata_copied':False,'removal_counts':removed}
        (stage/'receipt.json').write_text(json.dumps(receipt,indent=2),encoding='utf-8')
        os.chmod(stage/'receipt.json',0o600)
        require(not output.exists(),'output_exists')
        stage.rename(output)
        return {k:receipt[k] for k in ('status','code','cloud_release','network_denied','integrity_passed')}
    finally:
        if stage.exists():shutil.rmtree(stage)


def result_rows(path, cols):
    try:
        data=read_bytes(path,MAX_FILE).decode('utf-8-sig')
        rows=list(csv.reader(io.StringIO(data),strict=True))
    except (UnicodeError,csv.Error):
        raise Block('result_mismatch')
    names=[c['output'] for c in cols]
    require(rows and rows[0]==names and len(rows)<=MAX_ROWS+1,'result_mismatch')
    parsed=[];ids=set()
    for row in rows[1:]:
        require(len(row)==len(cols),'result_mismatch');record={}
        for value,c in zip(row,cols):
            name=c['output']
            if c['kind']=='id':
                require(re.fullmatch(r'p_[0-9a-f]{32}',value) is not None,'result_mismatch')
                require(value not in ids,'duplicate_identity');ids.add(value);record[name]=value
            elif c['kind']=='number':
                require(value!='' or c.get('nullable',True),'numeric_invalid')
                record[name]=None if value=='' else number(value,c)
            elif c['kind']=='category':
                require(value in c['mapping'].values(),'category_unknown');record[name]=value
        parsed.append(record)
    return parsed


def prepare_import(input_path, policy_path, output, job, results, import_policy):
    isolation_check()
    safe_path(job); safe_path(output)
    require(not output.exists() and not output.is_symlink(),'output_exists')
    p=load_policy(policy_path)
    try:
        q=json.loads(read_bytes(import_policy,4*1024*1024))
        require({'table','editable','fields'} <= set(q) <= {'table','editable','fields','derived','selection'},'policy_invalid')
        require(isinstance(q['editable'],list) and len(set(q['editable']))==len(q['editable']),'policy_invalid')
        require(isinstance(q['fields'],list) and 1<=len(q['fields'])<=MAX_COLS,'policy_invalid')
        specs=[t for t in p['tables'] if t['output']==q['table']]
        require(len(specs)==1,'policy_invalid');t=specs[0]
        cols=[c for c in t['columns'] if c['kind']!='drop'];idcol=next(c['output'] for c in cols if c['kind']=='id')
        editable={c['output'] for c in cols if c['kind'] in ('number','category')}
        require(set(q['editable'])<=editable,'policy_invalid')
        derived = q.get('derived', [])
        require(isinstance(derived,list) and len(derived)+len(cols)<=MAX_COLS,'policy_invalid')
        names = {c['output'] for c in cols}
        for c in derived:
            require(isinstance(c,dict) and c.get('kind') in ('number','category'),'policy_invalid')
            require(isinstance(c.get('output'),str) and TOKEN.fullmatch(c['output']) and c['output'] not in names,'policy_invalid')
            names.add(c['output'])
            if c['kind']=='number':
                require(set(c)=={'output','kind','min','max','nullable'},'policy_invalid')
                require(type(c['nullable']) is bool and type(c['min']) in (int,float) and type(c['max']) in (int,float)
                        and math.isfinite(c['min']) and math.isfinite(c['max']) and c['min']<=c['max'],'policy_invalid')
            else:
                require(set(c)=={'output','kind','mapping'},'policy_invalid')
                m=c['mapping']
                require(isinstance(m,dict) and 1<=len(m)<=50 and all(isinstance(k,str) and k.strip()
                        and isinstance(v,str) and TOKEN.fullmatch(v) for k,v in m.items())
                        and len(set(m.values()))==len(m),'policy_invalid')
        derived_names={c['output'] for c in derived}
        selection=q.get('selection',{'mode':'all'})
        require(isinstance(selection,dict),'policy_invalid')
        if selection.get('mode')=='all':
            require(set(selection)=={'mode'},'policy_invalid')
        else:
            require(set(selection)=={'mode','tokens'} and selection['mode']=='subset','policy_invalid')
            ts=selection['tokens']
            require(isinstance(ts,list) and 1<=len(ts)<=MAX_ROWS and all(isinstance(x,str) and re.fullmatch(r'p_[0-9a-f]{32}',x) for x in ts)
                    and len(set(ts))==len(ts),'policy_invalid')
        headers=[]
        for f in q['fields']:
            require(set(f)=={'source','column','header'} and isinstance(f['header'],str) and f['header'].strip(),'policy_invalid')
            headers.append(f['header'])
            if f['source']=='original':
                require(type(f['column']) is int and 0<=f['column']<len(t['columns']),'policy_invalid')
            elif f['source']=='result':
                require(f['column'] in editable|derived_names,'policy_invalid')
            else:raise Block('policy_invalid')
        require(len(headers)==len(set(headers)),'policy_invalid')
        binding=json.loads(read_bytes(job/'private/job.json',1024*1024))
        require(hashlib.sha256(read_bytes(input_path,MAX_FILE)).hexdigest()==binding['input_sha256'],'job_mismatch')
        require(hashlib.sha256(json.dumps(p,sort_keys=True).encode()).hexdigest()==binding['policy_sha256'],'job_mismatch')
        baseline=job/'private'/(t['output']+'.csv')
        require(hashlib.sha256(read_bytes(baseline,MAX_FILE)).hexdigest()==binding['tables'][t['output']],'job_mismatch')
        key=read_bytes(job/'private/job-key.bin',64);require(len(key)==32,'job_mismatch')
    except Block:raise
    except Exception:raise Block('policy_invalid')
    original,_,digest=input_tables(input_path,p)
    require(digest==binding['input_sha256'],'job_mismatch')
    source_rows=original[p['tables'].index(t)][1:]
    original_id_index=next(i for i,c in enumerate(t['columns']) if c['kind']=='id')
    def pseudonym(row):
        v=row[original_id_index];require(isinstance(v,str),'job_mismatch')
        return 'p_'+hmac.new(key,unicodedata.normalize('NFKC',v).strip().encode(),hashlib.sha256).hexdigest()[:32]
    raw_by_token={pseudonym(row):row for row in source_rows}
    base={r[idcol]:r for r in result_rows(baseline,cols)}
    edited={r[idcol]:r for r in result_rows(results,cols+derived)}
    require(set(base)==set(raw_by_token),'result_mismatch')
    allowed=set(base) if selection['mode']=='all' else set(selection['tokens'])
    require(allowed<=set(base) and set(edited)==allowed,'result_mismatch')
    for token,row in edited.items():
        for name in row:
            if name not in q['editable'] and name not in derived_names:require(row[name]==base[token][name],'unauthorized_change')
    # Restore in ORIGINAL source order, independent of model/result ordering.
    prepared=[]
    for raw in source_rows:
        if pseudonym(raw) not in allowed:continue
        row=edited[pseudonym(raw)];values=[]
        for f in q['fields']:
            value=raw[f['column']] if f['source']=='original' else row[f['column']]
            if f['source']=='result':
                dc=next((c for c in derived if c['output']==f['column'] and c['kind']=='category'),None)
                if dc is not None:value={v:k for k,v in dc['mapping'].items()}[value]
            if isinstance(value,Decimal):
                # Reject precision loss instead of quietly altering import values.
                as_float=float(value)
                require(math.isfinite(as_float) and Decimal(str(as_float))==value,'numeric_invalid')
                value=as_float
            values.append(value)
        prepared.append(values)
    require(output.parent.is_dir() and not output.parent.is_symlink(),'input_rejected')
    stage=Path(tempfile.mkdtemp(prefix='.privacy-',dir=output.parent));os.chmod(stage,0o700)
    try:
        import openpyxl
        private=stage/'private';private.mkdir(mode=0o700)
        wb=openpyxl.Workbook();ws=wb.active;ws.title='import'
        for row in [headers]+prepared:
            ws.append(row)
            for cell in ws[ws.max_row]:
                if isinstance(cell.value,str):cell.data_type='s'
        wb.properties.creator=None;wb.properties.lastModifiedBy=None
        wb.properties.created=datetime(2000,1,1);wb.properties.modified=datetime(2000,1,1)
        dest=private/'ready-to-import.xlsx';wb.save(dest);wb.close();os.chmod(dest,0o600)
        with ZipFile(dest) as z:parts={name:z.read(name) for name in z.namelist()}
        root=ET.fromstring(parts['docProps/core.xml'])
        for element in list(root):
            if element.tag.rsplit('}',1)[-1] in ['creator','lastModifiedBy','created','modified']:root.remove(element)
        parts['docProps/core.xml']=ET.tostring(root,encoding='utf-8',xml_declaration=True)
        with ZipFile(dest,'w',ZIP_DEFLATED) as z:
            for name,blob in parts.items():
                info=ZipInfo(name,(2000,1,1,0,0,0));info.compress_type=ZIP_DEFLATED;z.writestr(info,blob)
        # Verify every output cell, including leading zero IDs, before publishing.
        check=openpyxl.load_workbook(dest,data_only=False)
        actual=[list(r) for r in check.active.iter_rows(values_only=True)];check.close()
        require(actual==[headers]+prepared,'integrity_failed')
        def clip_safe(v):
            if v is None:return True
            if isinstance(v,(int,float)):return True
            return isinstance(v,str) and not any(x in v for x in '\t\r\n\x00') and not v.lstrip().startswith(('=','+','-','@'))
        clipboard_ready=all(clip_safe(v) for row in [headers]+prepared for v in row)
        if clipboard_ready:
            with (private/'ready-to-paste.tsv').open('w',encoding='utf-8',newline='') as f:
                writer=csv.writer(f,delimiter='\t',lineterminator='\n');writer.writerows([headers]+prepared)
            os.chmod(private/'ready-to-paste.tsv',0o600)
        receipt={'status':'local_ready','code':'local_ready','cloud_release':False,'network_denied':True,
                 'integrity_passed':True,'identity_restored_locally':True,
                 'system_submitted':False,'clipboard_written':False,'paste_file_ready':clipboard_ready,
                 'warning':'private_outputs_contain_original_identity'}
        receipt.update(exported_rows=len(prepared), omitted_rows=len(base)-len(prepared),
                       selection_mode=selection['mode'], derived_columns=len(derived))
        (stage/'receipt.json').write_text(json.dumps(receipt,indent=2),encoding='utf-8');os.chmod(stage/'receipt.json',0o600)
        require(not output.exists(),'output_exists');stage.rename(output)
        return {k:receipt[k] for k in ('status','code','cloud_release','network_denied','integrity_passed')}
    finally:
        if stage.exists():shutil.rmtree(stage)


def prepare_summary(input_path, policy_path, output, job, summary_policy):
    isolation_check()  # Before any file is opened.
    require(not output.exists() and not output.is_symlink(), 'output_exists')
    safe_path(output)
    require(not job.is_symlink() and not (job/'private').is_symlink(), 'input_rejected')
    p = load_policy(policy_path)
    try:
        q = json.loads(read_bytes(summary_policy, 16384))
        require(isinstance(q, dict) and set(q) == {'table','numeric','category'}, 'policy_invalid')
        require(isinstance(q['table'], str), 'policy_invalid')
        require(isinstance(q['numeric'], list) and len(q['numeric']) <= 3
                and all(isinstance(x,str) for x in q['numeric'])
                and len(set(q['numeric'])) == len(q['numeric']), 'policy_invalid')
        require(q['category'] is None or isinstance(q['category'], str), 'policy_invalid')
        specs = [t for t in p['tables'] if t['output'] == q['table']]
        require(len(specs) == 1, 'policy_invalid'); t = specs[0]
        cols = [c for c in t['columns'] if c['kind'] != 'drop']
        numbers = {c['output']:c for c in cols if c['kind'] == 'number'}
        categories = {c['output']:c for c in cols if c['kind'] == 'category'}
        require(set(q['numeric']) <= set(numbers), 'policy_invalid')
        require(q['category'] is None or q['category'] in categories, 'policy_invalid')
        binding = json.loads(read_bytes(job/'private/job.json', 1024*1024))
        require(hashlib.sha256(read_bytes(input_path, MAX_FILE)).hexdigest() == binding['input_sha256'], 'job_mismatch')
        require(hashlib.sha256(json.dumps(p, sort_keys=True).encode()).hexdigest() == binding['policy_sha256'], 'job_mismatch')
        base = job/'private'/(t['output']+'.csv')
        digest = hashlib.sha256(read_bytes(base, MAX_FILE)).hexdigest()
        require(digest == binding['tables'][t['output']], 'job_mismatch')
        rows = result_rows(base, cols)
        require(hashlib.sha256(read_bytes(base, MAX_FILE)).hexdigest() == digest, 'job_mismatch')
        identifier = next(c['output'] for c in cols if c['kind'] == 'id')
        require(len({r[identifier] for r in rows}) == len(rows), 'job_mismatch')
    except Block: raise
    except Exception: raise Block('policy_invalid')
    # Fixed threshold, no arbitrary filters, subgroup means, free text or extrema.
    threshold = max(10, t['min_group'])
    def cell(n): require(n == 0 or n >= threshold, 'summary_small_cell')
    require(len(rows) >= threshold, 'summary_small_cell')
    summary = {'version':1, 'row_count':len(rows), 'numeric':{}, 'category_counts':{}}
    mapping = []
    for i, name in enumerate(q['numeric'], 1):
        values = [r[name] for r in rows if r[name] is not None]
        n = len(values); missing = len(rows)-n
        cell(n); cell(missing)
        alias = 'measure_'+str(i)
        mean = str((sum(values, Decimal(0))/n).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)) if n else None
        summary['numeric'][alias] = {'n':n, 'missing':missing, 'mean_rounded':mean}
        mapping.append(alias+' 对应列位置 '+str(t['columns'].index(numbers[name])+1))
    if q['category'] is not None:
        c = categories[q['category']]; counts = Counter(r[c['output']] for r in rows)
        # Reject the WHOLE candidate if any cell or its complement is small.
        # Suppressing one cell alongside totals would reveal it by subtraction.
        for n in counts.values(): cell(n); cell(len(rows)-n)
        for i, value in enumerate(c['mapping'].values(), 1):
            alias = 'category_'+str(i)
            summary['category_counts'][alias] = counts.get(value, 0)
            mapping.append(alias+' 对应本地分类策略第 '+str(i)+' 项')
    require(output.parent.is_dir() and not output.parent.is_symlink(), 'input_rejected')
    stage = Path(tempfile.mkdtemp(prefix='.privacy-summary-', dir=output.parent)); os.chmod(stage,0o700)
    try:
        private = stage/'private'; private.mkdir(mode=0o700)
        blob = json.dumps(summary,ensure_ascii=False,indent=2).encode('utf-8')
        (private/'summary-review.json').write_bytes(blob)
        (private/'summary-review.txt').write_text(
            '本地待审阅汇总：尚未授权交给模型\n\n'+blob.decode()+'\n\n本地字段定位（不要随文件提供给模型）\n'
            +'\n'.join(mapping)+'\n\n数字已作程序检查，仍须判断业务许可、敏感含义及与其他结果组合后的识别风险。'
            '通过后由本人明确提供且仅提供 summary-review.json，并说明允许的用途与接收服务。'
            '不得提供整个 private 目录。均值保留两位小数；原始精度留在本地。\n',encoding='utf-8')
        receipt = {'status':'local_ready','code':'summary_review_required','cloud_release':False,
                   'network_denied':True,'integrity_passed':True}
        (stage/'receipt.json').write_text(json.dumps(receipt),encoding='utf-8')
        for f in stage.rglob('*'):
            if f.is_file():os.chmod(f,0o600)
        require(not output.exists(),'output_exists'); stage.rename(output)
        return receipt
    finally:
        if stage.exists():shutil.rmtree(stage)


def main():
    os.umask(0o077)
    try:
        parser=argparse.ArgumentParser(add_help=False)
        parser.add_argument('--input',required=True);parser.add_argument('--policy',required=True);parser.add_argument('--out',required=True)
        parser.add_argument('--mode',choices=['deidentify','prepare-import','prepare-summary'],default='deidentify')
        parser.add_argument('--job');parser.add_argument('--results');parser.add_argument('--import-policy');parser.add_argument('--summary-policy')
        args=parser.parse_args()
        if args.mode=='prepare-summary':
            result=prepare_summary(Path(args.input),Path(args.policy),Path(args.out),Path(args.job),Path(args.summary_policy))
        elif args.mode=='prepare-import':
            result=prepare_import(Path(args.input),Path(args.policy),Path(args.out),Path(args.job),Path(args.results),Path(args.import_policy))
        else:result=process(Path(args.input),Path(args.policy),Path(args.out))
        print(json.dumps(result));return 0
    except Block as e:
        print(json.dumps({'status':'blocked','code':str(e),'cloud_release':False}));return 2
    except BaseException:
        print(json.dumps({'status':'blocked','code':'processing_failed','cloud_release':False}));return 2


if __name__=='__main__':
    sys.exit(main())
